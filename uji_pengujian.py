import base64
import hashlib
import io
import math
import os
import secrets
import tempfile
import time
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageDraw
from cryptography.hazmat.primitives.ciphers.aead import (
    AESGCM,
    ChaCha20Poly1305,
)
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from crypto_utils import (
    MAGIC_AES,
    MAGIC_CHACHA,
    NONCE_SIZE,
    SALT_SIZE,
    decrypt_file,
    encrypt_file,
)


OUTPUT_FILE = Path("hasil_pengujian.xlsx")
ALGORITHMS = {
    "AES-256-GCM": (AESGCM, MAGIC_AES),
    "ChaCha20-Poly1305": (ChaCha20Poly1305, MAGIC_CHACHA),
}
PASSWORD = "PasswordPengujian123!"
BENCHMARK_PASSWORD = "BenchmarkPassword123"
BENCHMARK_FILE_SIZES = [
    ("1 KB", 1024),
    ("1 MB", 1024 * 1024),
    ("10 MB", 10 * 1024 * 1024),
]
BENCHMARK_REPEATS = 3
TAG_SIZE = 16


def make_pdf() -> bytes:
    content = b"BT /F1 18 Tf 72 720 Td (Sample PDF test document) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        (
            f"<< /Length {len(content)} >>\nstream\n".encode("ascii")
            + content
            + b"\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    document = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(document))
        document.extend(f"{object_number} 0 obj\n".encode("ascii"))
        document.extend(body)
        document.extend(b"\nendobj\n")

    xref_offset = len(document)
    document.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    document.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    document.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(document)


def make_test_files(folder: Path) -> list[Path]:
    histogram_phrase = (
        "CipherSafe file encryption test. Decrypted data must match the original.\n"
    ).encode("ascii")
    repeated_text = (
        histogram_phrase * ((1024 + len(histogram_phrase) - 1) // len(histogram_phrase))
    )[:1024]

    files = {
        "teks.txt": repeated_text,
        "data.csv": b"nama,nilai\nAyu,90\nBima,85\n",
        "data.json": b'{"app":"CipherSafe","active":true,"count":3}',
        "halaman.html": b"<!doctype html><title>Test</title><p>Sample</p>",
        "dokumen.pdf": make_pdf(),
        "arsip.zip": b"",
        "data.bin": bytes(range(256)) * 8,
    }

    for filename, data in files.items():
        (folder / filename).write_bytes(data)

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("readme.txt", "Sample ZIP test file")
    (folder / "arsip.zip").write_bytes(zip_buffer.getvalue())

    source_image = Path("data_uji/gambar.png")
    if source_image.exists():
        image = Image.open(source_image).convert("RGB")
    else:
        image = Image.new("RGB", (96, 64), "white")
        drawing = ImageDraw.Draw(image)
        drawing.rectangle((8, 8, 88, 56), fill=(42, 112, 168))
        drawing.ellipse((28, 16, 68, 56), fill=(230, 170, 55))

    image.save(folder / "gambar.png", format="PNG")
    image.save(folder / "gambar.jpg", format="JPEG", quality=90)
    image.save(folder / "gambar.gif", format="GIF")

    return sorted(folder.iterdir())


def byte_entropy(data: bytes) -> float:
    if not data:
        return 0.0

    total = len(data)
    return -sum(
        (count / total) * math.log2(count / total)
        for count in Counter(data).values()
    )


def raw_ciphertext(encoded: str, algorithm: str) -> bytes:
    payload = base64.b64decode(encoded, validate=True)
    magic = ALGORITHMS[algorithm][1]
    offset = len(magic) + SALT_SIZE + NONCE_SIZE
    return payload[offset:-TAG_SIZE]


def write_workbook(
    path: Path,
    file_rows: list[dict],
    entropy_rows: list[dict],
    avalanche_rows: list[dict],
    benchmark_rows: list[dict],
    distributions: list[tuple[str, bytes]],
) -> None:
    workbook = Workbook()
    file_sheet = workbook.active
    file_sheet.title = "Keberhasilan Dekripsi"
    benchmark_sheet = workbook.create_sheet("Waktu Enkripsi-Dekripsi")
    avalanche_sheet = workbook.create_sheet("Avalanche Effect")
    entropy_sheet = workbook.create_sheet("Entropi Ciphertext")
    histogram_sheet = workbook.create_sheet("Histogram Byte")

    def fill_sheet(sheet, rows: list[dict]) -> None:
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row[header] for header in headers])

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315C54")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False

        for column in sheet.columns:
            column_letter = column[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column_letter].width = min(max_length + 2, 48)

    fill_sheet(file_sheet, file_rows)
    fill_sheet(entropy_sheet, entropy_rows)
    fill_sheet(avalanche_sheet, avalanche_rows)
    fill_sheet(benchmark_sheet, benchmark_rows)

    counters = [Counter(distribution[1]) for distribution in distributions]
    plaintext_entropy = byte_entropy(distributions[0][1])
    aes_ciphertext = next(data for name, data in distributions if name == "AES-256-GCM")
    histogram_sheet.append(
        ["Nilai Byte"] + [distribution[0] for distribution in distributions]
    )
    for value in range(256):
        histogram_sheet.append(
            [value] + [counts.get(value, 0) for counts in counters]
        )
    for cell in histogram_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315C54")
    histogram_sheet.freeze_panes = "A2"
    histogram_sheet.auto_filter.ref = f"A1:{histogram_sheet.cell(1, len(distributions) + 1).column_letter}257"
    histogram_sheet.sheet_view.showGridLines = False
    histogram_sheet.column_dimensions["A"].width = 14
    for column_index in range(2, len(distributions) + 2):
        histogram_sheet.column_dimensions[
            histogram_sheet.cell(1, column_index).column_letter
        ].width = 24

    categories = Reference(histogram_sheet, min_col=1, min_row=2, max_row=257)

    def make_histogram(columns: list[tuple[int, str]], title: str) -> BarChart:
        chart = BarChart()
        chart.type = "col"
        chart.style = 2
        chart.title = title
        chart.y_axis.title = "Frekuensi"
        chart.x_axis.title = "Nilai byte"
        chart.height = 8.5
        chart.width = 12.5
        chart.gapWidth = 35

        for column, color in columns:
            data = Reference(histogram_sheet, min_col=column, min_row=1, max_row=257)
            chart.add_data(data, titles_from_data=True)
            series = chart.series[-1]
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color

        chart.set_categories(categories)
        chart.x_axis.tickLblSkip = 50
        if len(columns) == 1:
            chart.legend = None
        else:
            chart.legend.position = "b"
        return chart

    histogram_sheet.add_chart(
        make_histogram(
            [(2, "4C87B9")],
            f"Histogram Plaintext (entropi {plaintext_entropy:.2f})",
        ),
        "F2",
    )
    histogram_sheet.add_chart(
        make_histogram(
            [(3, "C8515B")],
            f"Histogram Ciphertext (entropi {byte_entropy(aes_ciphertext):.2f})",
        ),
        "M2",
    )

    try:
        workbook.save(path)
    except PermissionError as exc:
        raise SystemExit(
            f"Tutup {path.name} di Excel, lalu jalankan skrip kembali."
        ) from exc


def measure_key_avalanche() -> list[dict]:
    sample = (
        b"CipherSafe controlled key-avalanche test. "
        b"This sample is not application or user data. "
    ) * 32
    rows = []

    for algorithm in ALGORITHMS:
        cipher_type = ALGORITHMS[algorithm][0]
        base_key = os.urandom(32)
        differences = []

        for byte_index in range(len(base_key)):
            for bit_index in range(8):
                modified_key = bytearray(base_key)
                modified_key[byte_index] ^= 1 << bit_index
                nonce = os.urandom(12)

                baseline = cipher_type(base_key).encrypt(
                    nonce, sample, None
                )[:-TAG_SIZE]
                modified = cipher_type(bytes(modified_key)).encrypt(
                    nonce, sample, None
                )[:-TAG_SIZE]
                changed_bits = sum(
                    (left_byte ^ right_byte).bit_count()
                    for left_byte, right_byte in zip(baseline, modified)
                )
                differences.append(changed_bits / (len(sample) * 8) * 100)

        rows.append({
            "Algoritma": algorithm,
            "Jenis Perubahan": "Satu bit kunci diubah",
            "Bagian yang Diukur": "Ciphertext saja; tag autentikasi dikecualikan",
            "Jumlah Pengujian": len(differences),
            "Rata-rata Bit Berubah (%)": round(sum(differences) / len(differences), 4),
            "Minimum (%)": round(min(differences), 4),
            "Maksimum (%)": round(max(differences), 4),
        })

    return rows


def run_file_tests(work_dir: Path) -> tuple[list[dict], list[dict], list[tuple[str, bytes]]]:
    test_files = make_test_files(work_dir / "inputs")
    file_rows = []
    entropy_rows = []
    histogram_data = []
    text_sample = next(path for path in test_files if path.name == "teks.txt")

    for input_path in test_files:
        original_data = input_path.read_bytes()
        original_hash = hashlib.sha256(original_data).hexdigest()

        for algorithm in ALGORITHMS:
            encrypted_path = work_dir / f"{input_path.stem}_{algorithm}.utsk"
            decrypted_path = work_dir / f"{input_path.stem}_{algorithm}_decrypted"
            error = ""
            ciphertext_data = None
            decrypted_hash = ""
            status = "FAIL"

            try:
                encrypt_file(
                    str(input_path),
                    str(encrypted_path),
                    PASSWORD,
                    algorithm,
                )
                decrypt_file(
                    str(encrypted_path),
                    str(decrypted_path),
                    PASSWORD,
                )
                decrypted_data = decrypted_path.read_bytes()
                decrypted_hash = hashlib.sha256(decrypted_data).hexdigest()
                status = "PASS" if decrypted_data == original_data else "FAIL"
                if status == "FAIL":
                    error = "Hasil dekripsi berbeda dari file asli."

                encoded = encrypted_path.read_text(encoding="utf-8")
                ciphertext_data = raw_ciphertext(encoded, algorithm)
            except Exception as exc:
                error = f"{type(exc).__name__}: {exc}"

            file_rows.append({
                "Nama File": input_path.name,
                "Format": input_path.suffix.lstrip(".").upper(),
                "Ukuran (byte)": len(original_data),
                "Algoritma": algorithm,
                "Status": status,
                "SHA-256 Asli": original_hash,
                "SHA-256 Hasil Dekripsi": decrypted_hash,
                "Catatan": error,
            })

            if ciphertext_data is not None:
                entropy_rows.append({
                    "Nama File": input_path.name,
                    "Algoritma": algorithm,
                    "Entropi Asli (bit/byte)": round(byte_entropy(original_data), 6),
                    "Entropi Ciphertext (bit/byte)": round(byte_entropy(ciphertext_data), 6),
                    "Ukuran Ciphertext (byte)": len(ciphertext_data),
                })
                if input_path == text_sample and algorithm == "AES-256-GCM":
                    histogram_data.append((
                        algorithm,
                        ciphertext_data,
                    ))

    histogram_data.insert(0, (
        "Teks asli",
        text_sample.read_bytes(),
    ))
    return file_rows, entropy_rows, histogram_data


def measure_benchmarks(work_dir: Path) -> list[dict]:
    benchmark_dir = work_dir / "benchmark"
    benchmark_dir.mkdir()
    results = []

    for size_name, size_bytes in BENCHMARK_FILE_SIZES:
        original_file = benchmark_dir / f"input_{size_name.replace(' ', '')}.bin"
        original_file.write_bytes(secrets.token_bytes(size_bytes))

        for algorithm in ALGORITHMS:
            encrypt_times = []
            decrypt_times = []

            for run_number in range(1, BENCHMARK_REPEATS + 1):
                encrypted_file = benchmark_dir / (
                    f"encrypted_{size_name.replace(' ', '')}_"
                    f"{algorithm}_{run_number}.utsk"
                )
                decrypted_file = benchmark_dir / (
                    f"decrypted_{size_name.replace(' ', '')}_"
                    f"{algorithm}_{run_number}.bin"
                )

                start = time.perf_counter()
                encrypt_file(
                    str(original_file),
                    str(encrypted_file),
                    BENCHMARK_PASSWORD,
                    algorithm,
                )
                encrypt_times.append(time.perf_counter() - start)

                start = time.perf_counter()
                decrypt_file(
                    str(encrypted_file),
                    str(decrypted_file),
                    BENCHMARK_PASSWORD,
                )
                decrypt_times.append(time.perf_counter() - start)

                if original_file.read_bytes() != decrypted_file.read_bytes():
                    raise ValueError(
                        f"Hasil benchmark tidak cocok: {decrypted_file}"
                    )

            results.append({
                "Ukuran File": size_name,
                "Ukuran (byte)": size_bytes,
                "Algoritma": algorithm,
                "Rata-rata Enkripsi (detik)": sum(encrypt_times) / BENCHMARK_REPEATS,
                "Rata-rata Dekripsi (detik)": sum(decrypt_times) / BENCHMARK_REPEATS,
                "Jumlah Pengulangan": BENCHMARK_REPEATS,
            })

    return results


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="utski-tests-") as temporary:
        work_dir = Path(temporary)
        (work_dir / "inputs").mkdir()
        file_rows, entropy_rows, histogram_data = run_file_tests(work_dir)
        benchmark_rows = measure_benchmarks(work_dir)

    avalanche_rows = measure_key_avalanche()
    write_workbook(
        OUTPUT_FILE,
        file_rows,
        entropy_rows,
        avalanche_rows,
        benchmark_rows,
        histogram_data,
    )

    passed = sum(row["Status"] == "PASS" for row in file_rows)
    failed = len(file_rows) - passed
    print(f"Uji file: {passed}/{len(file_rows)} berhasil; {failed} gagal.")
    print(f"Hasil pengujian dan analisis: {OUTPUT_FILE.resolve()}")
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()